"""Text-Extraktion aus Dokumenten (PDF, PPTX, DOCX, HTML, TXT, Untertitel).

Konvention: Die Extraktion erzeugt Markdown mit Struktur-Ankern
("## Seite 3", "## Slide 7"), damit spaeter zitiert und gechunkt werden kann.
"""
from __future__ import annotations

import html
import re
from pathlib import Path


class ExtractionError(RuntimeError):
    pass


def _require(module: str, package: str):
    try:
        return __import__(module)
    except ImportError as exc:  # pragma: no cover
        raise ExtractionError(f"Paket fehlt: `pip install {package}`") from exc


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\u00ad", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------------------------------------------------------------- PDF
def from_pdf(path: Path) -> tuple[str, dict]:
    _require("pypdf", "pypdf")
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("PDF ist passwortgeschuetzt.") from exc

    pages: list[str] = []
    empty = 0
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if not text.strip():
            empty += 1
            continue
        pages.append(f"## Seite {i}\n\n{_clean(text)}")

    if not pages:
        raise ExtractionError(
            "Kein extrahierbarer Text – vermutlich ein Scan. "
            "Loesung: PDF vorher per OCR aufbereiten (z. B. `ocrmypdf in.pdf out.pdf`)."
        )
    return "\n\n".join(pages), {"pages": len(reader.pages), "pages_without_text": empty}


# ---------------------------------------------------------------- PowerPoint
def from_pptx(path: Path) -> tuple[str, dict]:
    _require("pptx", "python-pptx")
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append(f"_Notizen:_ {notes}")   # Dozentennotizen sind oft das Gold
        if lines:
            slides.append(f"## Slide {i}\n\n" + _clean("\n".join(lines)))
    if not slides:
        raise ExtractionError("Keine Textinhalte in der Praesentation gefunden (nur Bilder?).")
    return "\n\n".join(slides), {"slides": len(prs.slides)}


# ---------------------------------------------------------------- Word / HTML / Plain
def from_docx(path: Path) -> tuple[str, dict]:
    _require("docx", "python-docx")
    import docx

    doc = docx.Document(str(path))
    blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            blocks.append(" | ".join(c.text.strip() for c in row.cells))
    if not blocks:
        raise ExtractionError("Dokument enthaelt keinen Text.")
    return _clean("\n\n".join(blocks)), {"paragraphs": len(blocks)}


def from_html(path: Path) -> tuple[str, dict]:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise ExtractionError("Paket fehlt: `pip install beautifulsoup4`") from exc
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return _clean(soup.get_text("\n")), {}


def from_plain(path: Path) -> tuple[str, dict]:
    return _clean(path.read_text(encoding="utf-8", errors="ignore")), {}


_TS_RE = re.compile(r"^\d{1,2}:\d{2}:\d{2}[.,]\d{3}\s-->")


def from_subtitles(path: Path) -> tuple[str, dict]:
    """VTT/SRT: Zeitmarken zu Ankern verdichten, Duplikate entfernen."""
    out: list[str] = []
    stamp = ""
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        line = line.strip()
        if not line or line.isdigit() or line.upper().startswith("WEBVTT"):
            continue
        if _TS_RE.match(line):
            stamp = line.split(" --> ")[0].split(".")[0].split(",")[0]
            continue
        text = html.unescape(re.sub(r"<[^>]+>", "", line))
        if out and out[-1].endswith(text):
            continue
        if stamp and (not out or len(out[-1]) > 400):
            out.append(f"\n## {stamp}\n")
            stamp = ""
        out.append(text)
    return _clean(" ".join(out)), {"format": path.suffix.lstrip(".")}


_DISPATCH = {
    ".pdf": (from_pdf, "slides_or_doc"),
    ".pptx": (from_pptx, "slides"),
    ".docx": (from_docx, "document"),
    ".txt": (from_plain, "document"),
    ".md": (from_plain, "document"),
    ".csv": (from_plain, "document"),
    ".html": (from_html, "document"),
    ".htm": (from_html, "document"),
    ".vtt": (from_subtitles, "transcript"),
    ".srt": (from_subtitles, "transcript"),
}


def extract(path: Path) -> tuple[str, str, dict]:
    """-> (text, kind, meta). Wirft ExtractionError bei nicht unterstuetztem Typ."""
    suffix = path.suffix.lower()
    if suffix not in _DISPATCH:
        raise ExtractionError(f"Dateityp {suffix} wird nicht unterstuetzt.")
    func, kind = _DISPATCH[suffix]
    text, meta = func(path)
    if kind == "slides_or_doc":
        kind = "slides" if "Slide" in text[:2000] else "document"
    return text, kind, meta

